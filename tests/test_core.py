import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd

import sys

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import word_freq_analyzer as analyzer
from llm_sentence_analyzer import LLMAnalyzerConfig, QwenSentenceAnalyzer
from wordfreq_app import settings as app_settings
from wordfreq_app.version import version_tuple


class CoreRegressionTests(unittest.TestCase):
    def test_parse_mixed_year_formats_without_dt_failure(self):
        values = pd.Series([
            "2024",
            "2024-03-15",
            "2011/1/1-2011/12/31",
            "2009年",
            45292,  # Excel serial date: 2024-01-01
            "无法识别",
        ])
        self.assertEqual(
            analyzer.parse_year_column(values).tolist(),
            [2024, 2024, 2011, 2009, 2024, 0],
        )
        self.assertEqual(
            analyzer.parse_month_column(pd.Series([
                "2024", "2024-03-15", "2024年4月", "2011/1/1-2011/12/31", 45292,
            ])).tolist(),
            [0, 3, 4, 1, 1],
        )

    def test_excel_headers_are_consistent_with_preview(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "blank_header.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["companyid", None, "调研报告内容"])
            ws.append([1, "2024", "公司积极保护生态环境并修复湿地。"])
            wb.save(path)
            wb.close()

            columns = analyzer._read_columns_fast(str(path))
            frame = analyzer.read_data_file(str(path))
            self.assertEqual(columns, list(frame.columns))
            self.assertEqual(columns[1], "Unnamed: 1")

    def test_dictionary_excel_header_is_not_imported_as_keyword(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "dict.xlsx"
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["分类", "关键词"])
            ws.append(["环境", "生态环境"])
            wb.save(path)
            wb.close()

            manager = analyzer.DictionaryManager()
            manager.import_file(str(path))
            self.assertEqual(manager.data, {"环境": ["生态环境"]})

    def test_realistic_xlsx_pipeline_writes_main_and_sentence_outputs(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "调研内容.xlsx"
            output = tmp_path / "result.xlsx"
            pd.DataFrame([
                {"companyid": "000001", "公司简称": "甲公司", "日期": "2024年",
                 "调研报告内容": "公司积极保护生态环境并修复湿地。"},
                {"companyid": "2", "公司简称": "乙公司", "日期": "2024-03-15",
                 "调研报告内容": "公司持续推进绿色发展和环境治理。"},
                {"companyid": "1", "公司简称": "甲公司", "日期": "2024-04-20",
                 "调研报告内容": "公司继续推进绿色发展。"},
            ]).to_excel(source, index=False)

            manager = analyzer.DictionaryManager()
            manager.add_category("环境")
            manager.add_word("环境", "生态环境")
            manager.add_word("环境", "绿色发展")

            analyzer.run_analysis(
                files=[str(source)],
                dict_mgr=manager,
                col_stkcd="companyid",
                col_year="日期",
                text_columns=["调研报告内容"],
                output_path=str(output),
                use_regex=True,
                export_sentences=True,
            )

            self.assertTrue(output.exists())
            sentence_output = tmp_path / "result_sentences.xlsx"
            self.assertTrue(sentence_output.exists())
            with pd.ExcelFile(output) as workbook:
                self.assertEqual(
                    workbook.sheet_names,
                    [
                        "原始记录统计", "原始记录关键词", "分类汇总", "数据质量",
                        "异常记录", "输入文件清单", "运行元数据", "词典诊断", "分析说明",
                    ],
                )
            panel = pd.read_excel(output, sheet_name="原始记录统计", dtype=str)
            self.assertEqual(set(panel["公司代码"]), {"000001", "000002"})
            self.assertEqual(len(panel), 3)
            self.assertEqual((panel["公司代码"] == "000001").sum(), 2)
            self.assertEqual(
                int(panel.loc[panel["公司代码"] == "000002", "月份"].iloc[0]), 3
            )
            self.assertIn("日期", panel.columns)
            self.assertIn("来源文件", panel.columns)
            self.assertIn("源文件行号", panel.columns)
            self.assertTrue((pd.to_numeric(panel["总计"], errors="coerce") > 0).all())
            sentences = pd.read_excel(sentence_output, sheet_name="命中句子")
            self.assertEqual(len(sentences), 3)
            self.assertFalse((tmp_path / "result_checkpoint").exists())

    def test_raw_mode_keeps_selected_source_columns(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "source.xlsx"
            output = tmp_path / "result.xlsx"
            pd.DataFrame([{
                "companyid": "1", "公司简称": "甲公司", "调研机构": "乙机构",
                "日期": "2024-03-01", "内容": "公司推进绿色发展。",
            }]).to_excel(source, index=False)
            manager = analyzer.DictionaryManager()
            manager.add_category("环境")
            manager.add_word("环境", "绿色发展")

            analyzer.run_analysis(
                files=[str(source)], dict_mgr=manager,
                col_stkcd="companyid", col_year="日期", text_columns=["内容"],
                retained_columns=["公司简称", "调研机构"],
                output_path=str(output), aggregation_mode="raw",
            )
            result = pd.read_excel(output, sheet_name="原始记录统计", dtype=str)
            self.assertEqual(result.loc[0, "公司简称"], "甲公司")
            self.assertEqual(result.loc[0, "调研机构"], "乙机构")

    def test_monthly_mode_aggregates_and_reports_unknown_month(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "source.xlsx"
            output = tmp_path / "result.xlsx"
            pd.DataFrame([
                {"code": "1", "日期": "2024-03-01", "内容": "公司持续推进绿色绿色发展工作。"},
                {"code": "1", "日期": "2024-03-20", "内容": "公司持续推进绿色发展工作。"},
                {"code": "1", "日期": "2024-04-01", "内容": "公司持续推进绿色发展工作。"},
                {"code": "1", "日期": "2024年", "内容": "公司持续推进绿色发展工作。"},
            ]).to_excel(source, index=False)
            manager = analyzer.DictionaryManager()
            manager.add_category("环境")
            manager.add_word("环境", "绿色")

            analyzer.run_analysis(
                files=[str(source)], dict_mgr=manager,
                col_stkcd="code", col_year="日期", text_columns=["内容"],
                output_path=str(output), aggregation_mode="monthly",
            )
            result = pd.read_excel(output, sheet_name="公司月份分类统计")
            self.assertEqual(result["月份"].tolist(), [3, 4])
            self.assertEqual(result["环境"].tolist(), [3, 1])
            quality = pd.read_excel(output, sheet_name="数据质量")
            total = quality.iloc[-1]
            self.assertEqual(int(total["有效记录数"]), 4)
            self.assertEqual(int(total["进入统计记录数"]), 3)
            self.assertEqual(int(total["月份无法识别"]), 1)

    def test_excel_layout_detects_nonfirst_sheet_and_header_row(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "layout.xlsx"
            workbook = openpyxl.Workbook()
            notes = workbook.active
            notes.title = "说明"
            notes.append(["本文件为调研数据说明"])
            data = workbook.create_sheet("数据")
            data.append(["调研数据汇总表"])
            data.append([])
            data.append(["companyid", "日期", "调研报告内容"])
            data.append(["1", "2024-05-01", "绿色发展"])
            workbook.save(path)
            workbook.close()

            self.assertEqual(
                analyzer._read_columns_fast(str(path)),
                ["companyid", "日期", "调研报告内容"],
            )
            frame = analyzer.read_data_file(str(path), nrows=10)
            self.assertEqual(frame.iloc[0]["调研报告内容"], "绿色发展")

    def test_shared_string_xlsx_headers_and_scan_progress_are_detected(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "shared_strings.xlsx"
            import xlsxwriter

            workbook = xlsxwriter.Workbook(path)
            sheet = workbook.add_worksheet("数据")
            sheet.write_row(0, 0, ["股票代码", "提问时间", "提问内容"])
            sheet.write_row(1, 0, ["000001", "2024-03-01", "绿色发展"])
            workbook.close()

            with zipfile.ZipFile(path) as archive:
                self.assertIn("xl/sharedStrings.xml", archive.namelist())
            self.assertEqual(
                analyzer._read_columns_fast(str(path)),
                ["股票代码", "提问时间", "提问内容"],
            )
            progress = []
            columns, frequency = analyzer.scan_all_columns(
                [str(path)], lambda done, total: progress.append((done, total)),
            )
            self.assertEqual(columns, ["股票代码", "提问时间", "提问内容"])
            self.assertEqual(frequency["股票代码"], 1)
            self.assertEqual(progress, [(1, 1)])

    def test_quality_report_records_invalid_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            tmp_path = Path(tmp)
            source = tmp_path / "source.xlsx"
            output = tmp_path / "result.xlsx"
            pd.DataFrame([
                {"code": "", "日期": "2024-01-01", "内容": "绿色"},
                {"code": "1", "日期": "无法识别", "内容": "绿色"},
                {"code": "2", "日期": "2024-03-01", "内容": ""},
            ]).to_excel(source, index=False)
            manager = analyzer.DictionaryManager()
            manager.add_category("环境")
            manager.add_word("环境", "绿色")
            analyzer.run_analysis(
                files=[str(source)], dict_mgr=manager,
                col_stkcd="code", col_year="日期", text_columns=["内容"],
                output_path=str(output), aggregation_mode="raw",
            )
            quality = pd.read_excel(output, sheet_name="数据质量").iloc[-1]
            self.assertEqual(int(quality["公司代码为空"]), 1)
            self.assertEqual(int(quality["年份无法识别"]), 1)
            self.assertEqual(int(quality["文本为空"]), 1)
            issues = pd.read_excel(output, sheet_name="异常记录")
            self.assertEqual(set(issues["问题类型"]), {"公司代码为空", "年份无法识别", "文本为空"})

    def test_settings_never_write_api_key_to_json(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "settings.json"
            with patch.object(app_settings, "settings_path", return_value=path):
                app_settings.save_settings({"api_key": "secret", "aggregation_mode": "raw"})
                raw = path.read_text(encoding="utf-8")
                self.assertNotIn("secret", raw)
                self.assertEqual(app_settings.load_settings()["aggregation_mode"], "raw")

    def test_version_tuple_handles_release_tags(self):
        self.assertGreater(version_tuple("v4.2.0"), version_tuple("4.1"))

    def test_llm_partial_json_is_not_reported_as_success(self):
        model = QwenSentenceAnalyzer(LLMAnalyzerConfig.from_inputs(api_key="test-key"))
        partial = model._normalize_result({"rel": 1, "time": 2})
        self.assertEqual(partial["LLM分析状态"], "失败")
        complete = model._normalize_result({
            "rel": 1, "time": 2, "voice": 0, "type": 2,
            "cert": 2, "quant": 1, "tone": 1, "conf": 2,
        })
        self.assertEqual(complete["LLM分析状态"], "成功")


if __name__ == "__main__":
    unittest.main()
