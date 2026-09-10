"""Automatic worksheet and header-row detection for spreadsheet inputs."""

from __future__ import annotations

import os
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable

HEADER_SCAN_ROWS = 20
HEADER_HINTS = (
    "公司", "企业", "股票", "证券", "代码", "编号", "日期", "时间", "年份",
    "year", "date", "code", "company", "文本", "内容", "正文", "回复",
    "问题", "标题", "摘要", "text", "content", "reply", "question", "body",
)


def _clean(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in {"nan", "none", "nat"} else text


def choose_header_row(rows: Iterable[Iterable[Any]]) -> tuple[int, float]:
    """Choose a zero-based header row and return a simple 0–1 confidence."""
    prepared = [list(row) for row in rows]
    if not prepared:
        return 0, 0.0

    scored: list[tuple[float, int]] = []
    max_width = max((len(row) for row in prepared), default=1)
    for index, row in enumerate(prepared[:HEADER_SCAN_ROWS]):
        values = [_clean(value) for value in row]
        nonempty = [value for value in values if value]
        if not nonempty:
            scored.append((-1000.0, index))
            continue
        unique_ratio = len(set(nonempty)) / len(nonempty)
        string_like = sum(not value.replace(".", "", 1).isdigit() for value in nonempty)
        hints = sum(
            1 for value in nonempty
            if any(hint in value.lower() for hint in HEADER_HINTS)
        )
        following_nonempty = 0
        if index + 1 < len(prepared):
            following_nonempty = sum(bool(_clean(v)) for v in prepared[index + 1])
        # A real header tends to be wide, unique, textual, followed by data, and
        # close to the top. Domain hints deliberately receive the largest weight.
        score = (
            len(nonempty) * 3.0
            + unique_ratio * 4.0
            + string_like * 1.5
            + hints * 15.0
            + min(following_nonempty, len(nonempty)) * 0.5
            + max(0, 5 - index) * 0.75
        )
        if len(nonempty) == 1:
            score -= 12.0
        scored.append((score, index))

    scored.sort(key=lambda item: (-item[0], item[1]))
    best_score, best_index = scored[0]
    theoretical = max(20.0, max_width * 6.0 + 15.0)
    confidence = max(0.0, min(1.0, best_score / theoretical))
    return best_index, round(confidence, 3)


@lru_cache(maxsize=512)
def _detect_xlsx_cached(path: str, size: int, mtime_ns: int) -> tuple[str, int, float]:
    del size, mtime_ns
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        candidates: list[tuple[float, int, str, int]] = []
        for sheet_index, worksheet in enumerate(workbook.worksheets):
            rows = list(worksheet.iter_rows(
                min_row=1, max_row=HEADER_SCAN_ROWS + 1, values_only=True,
            ))
            header_row, confidence = choose_header_row(rows)
            nonempty = sum(bool(_clean(value)) for value in rows[header_row]) if rows else 0
            candidates.append((confidence, nonempty, worksheet.title, header_row))
        if not candidates:
            return "", 0, 0.0
        candidates.sort(key=lambda item: (-item[0], -item[1]))
        confidence, _, sheet_name, header_row = candidates[0]
        return sheet_name, header_row, confidence
    finally:
        workbook.close()


def detect_xlsx_layout(path: str) -> tuple[str, int, float]:
    stat = os.stat(path)
    return _detect_xlsx_cached(os.path.abspath(path), stat.st_size, stat.st_mtime_ns)


def detect_xls_layout(path: str) -> tuple[str | int, int, float]:
    import pandas as pd

    workbook = pd.ExcelFile(path)
    candidates: list[tuple[float, int, str, int]] = []
    for sheet_name in workbook.sheet_names:
        preview = pd.read_excel(
            workbook, sheet_name=sheet_name, header=None, nrows=HEADER_SCAN_ROWS + 1,
        )
        rows = preview.where(preview.notna(), None).values.tolist()
        header_row, confidence = choose_header_row(rows)
        nonempty = sum(bool(_clean(value)) for value in rows[header_row]) if rows else 0
        candidates.append((confidence, nonempty, sheet_name, header_row))
    if not candidates:
        return 0, 0, 0.0
    candidates.sort(key=lambda item: (-item[0], -item[1]))
    confidence, _, sheet_name, header_row = candidates[0]
    return sheet_name, header_row, confidence


def describe_layout(path: str) -> tuple[str, int, float]:
    extension = Path(path).suffix.lower()
    if extension == ".xlsx":
        return detect_xlsx_layout(path)
    if extension == ".xls":
        sheet, row, confidence = detect_xls_layout(path)
        return str(sheet), row, confidence
    return "", 0, 1.0
